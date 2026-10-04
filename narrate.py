"""Turn ebooks (EPUB, PDF, TXT, DOCX) into narrated audiobooks with Kokoro TTS.

Examples:
    python narrate.py book.epub                    # -> output/<book>.m4b with chapters
    python narrate.py book.pdf --voice bm_george   # British male voice
    python narrate.py book.epub --play             # read aloud live through speakers
    python narrate.py book.epub --list-chapters
    python narrate.py book.epub --chapters 3-5 --format mp3
    python narrate.py --list-voices
"""
import argparse
import re
import shutil
import subprocess
import sys
import warnings
from dataclasses import dataclass
from pathlib import Path

warnings.filterwarnings("ignore")

SAMPLE_RATE = 24000

VOICES = {
    "American female": "af_heart af_bella af_nicole af_aoede af_kore af_sarah af_nova af_sky af_alloy af_jessica af_river",
    "American male": "am_michael am_fenrir am_puck am_echo am_eric am_liam am_onyx am_adam am_santa",
    "British female": "bf_emma bf_isabella bf_alice bf_lily",
    "British male": "bm_george bm_fable bm_lewis bm_daniel",
}


@dataclass
class Chapter:
    title: str
    text: str


# ---------------------------------------------------------------- extraction

def clean(text: str) -> str:
    text = text.replace("­", "")                       # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)            # re-join hyphenated line breaks
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{2,}", "\n\n", text)
    return text.strip()


def extract_epub(path: Path) -> list[Chapter]:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    toc_titles = {}

    def walk(entries):
        for e in entries:
            if isinstance(e, tuple):
                section, children = e
                if getattr(section, "href", None):
                    toc_titles.setdefault(section.href.split("#")[0], section.title)
                walk(children)
            elif hasattr(e, "href"):
                toc_titles.setdefault(e.href.split("#")[0], e.title)
    walk(book.toc)

    chapters = []
    for item_id, _ in book.spine:
        item = book.get_item_with_id(item_id)
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        soup = BeautifulSoup(item.get_content(), "lxml")
        for tag in soup(["script", "style", "nav", "aside", "sup"]):
            tag.decompose()
        for tag in soup.select('figcaption, [class*="caption"], [role="figure"], [class*="pageno"], [class*="pagenum"]'):
            tag.decompose()             # picture captions and printed page numbers aren't read aloud
        for br in soup("br"):
            br.replace_with(" ")
        for img in soup("img"):          # drop-cap letters are often images with the letter as alt text
            alt = (img.get("alt") or "").strip()
            img.replace_with(alt if len(alt) <= 2 else "")
        # join without a separator so styled drop caps ("<span>I</span>t") stay one word
        blocks = [(b.name, " ".join(b.get_text("").split()))
                  for b in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote", "div"])
                  if not b.find(["p", "div", "li", "blockquote"])]
        blocks = [(n, t) for n, t in blocks if t]

        # Many EPUBs put several chapters in one file; split at the top heading level that repeats.
        counts = {h: sum(1 for n, _ in blocks if n == h) for h in ("h1", "h2", "h3")}
        split_at = next((h for h in ("h1", "h2", "h3") if counts[h] >= 2), None)
        splitters = {"h1", "h2", "h3"}
        splitters = {h for h in splitters if split_at and h <= split_at}
        chunks, cur = [], []
        for n, t in blocks:
            if n in splitters and cur and any(m not in splitters for m, _ in cur):
                chunks.append(cur)
                cur = []
            cur.append((n, t))
        if cur:
            chunks.append(cur)

        for chunk in chunks:
            text = clean("\n\n".join(t for _, t in chunk))
            if len(text) < 40:
                continue
            heads = [i for i, (n, _) in enumerate(chunk) if n in ("h1", "h2", "h3")]
            title = None
            if len(chunks) == 1:
                title = toc_titles.get(item.get_name())
            if not title and heads:
                named = re.compile(r"^(chapter|part|book|prologue|epilogue|section|letter|stave)\b", re.I)
                i = next((i for i in heads if named.match(chunk[i][1])), heads[0])
                title = chunk[i][1]
                # "Chapter 3" followed directly by a heading with its name
                if i + 1 < len(chunk) and chunk[i + 1][0] in ("h1", "h2", "h3", "h4") and len(title) < 20:
                    title = f"{title} {chunk[i + 1][1]}"
            chapters.append(Chapter((title or f"Section {len(chapters) + 1}")[:100], text))
    return chapters


def extract_pdf(path: Path) -> list[Chapter]:
    import pymupdf

    doc = pymupdf.open(str(path))

    def page_text(i: int) -> str:
        # Join lines inside a block into one paragraph; blocks become paragraphs.
        paras = []
        for b in doc[i].get_text("blocks"):
            if b[6] != 0:      # skip images
                continue
            t = b[4].strip()
            if re.fullmatch(r"\d{1,4}", t):   # bare page numbers
                continue
            paras.append(re.sub(r"(\w)-\n(\w)", r"\1\2", t).replace("\n", " "))
        return "\n\n".join(paras)

    toc = [(lvl, title, page - 1) for lvl, title, page in doc.get_toc() if page >= 1]
    top = min((lvl for lvl, _, _ in toc), default=1)
    marks = [(t, p) for lvl, t, p in toc if lvl == top]
    chapters = []
    if marks:
        if marks[0][1] > 0:
            marks.insert(0, ("Front matter", 0))
        for n, (title, start) in enumerate(marks):
            end = marks[n + 1][1] if n + 1 < len(marks) else len(doc)
            text = clean("\n\n".join(page_text(i) for i in range(start, max(end, start + 1))))
            if len(text) >= 40:
                chapters.append(Chapter(title.strip()[:100], text))
    else:
        step = 10
        for start in range(0, len(doc), step):
            end = min(start + step, len(doc))
            text = clean("\n\n".join(page_text(i) for i in range(start, end)))
            if len(text) >= 40:
                chapters.append(Chapter(f"Pages {start + 1}-{end}", text))
    return chapters


def extract_docx(path: Path) -> list[Chapter]:
    import docx

    chapters, title, buf = [], "Start", []
    for p in docx.Document(str(path)).paragraphs:
        if p.style.name.lower().startswith("heading") and p.text.strip():
            if buf:
                chapters.append(Chapter(title, clean("\n\n".join(buf))))
            title, buf = p.text.strip()[:100], [p.text]
        elif p.text.strip():
            buf.append(p.text)
    if buf:
        chapters.append(Chapter(title, clean("\n\n".join(buf))))
    return chapters


def extract_txt(path: Path) -> list[Chapter]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?<!\n)\n(?!\n)", " ", raw)   # unwrap hard-wrapped lines
    heading = re.compile(r"^\s*(chapter|part|book|prologue|epilogue)\b[^\n]{0,80}$", re.I | re.M)
    cuts = [m.start() for m in heading.finditer(raw)]
    if not cuts:
        return [Chapter(path.stem, clean(raw))]
    if cuts[0] > 0 and raw[:cuts[0]].strip():
        cuts.insert(0, 0)
    chapters = []
    for n, s in enumerate(cuts):
        part = raw[s:cuts[n + 1] if n + 1 < len(cuts) else len(raw)]
        first = part.strip().split("\n", 1)[0]
        chapters.append(Chapter(first[:100] if s else "Opening", clean(part)))
    return chapters


def extract(path: Path) -> list[Chapter]:
    ext = path.suffix.lower()
    if ext == ".epub":
        return extract_epub(path)
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in (".txt", ".md", ".text"):
        return extract_txt(path)
    sys.exit(f"Unsupported file type: {ext} (supported: .epub .pdf .docx .txt .md)")


# ---------------------------------------------------------------- speech

class Narrator:
    REPO = "hexgrad/Kokoro-82M"

    def __init__(self, voice: str = "af_heart", speed: float = 1.0):
        import threading
        import torch
        from kokoro import KModel

        device = "cuda" if torch.cuda.is_available() else "cpu"
        print(f"Loading Kokoro on {device}...")
        self.model = KModel(repo_id=self.REPO).to(device).eval()
        self.pipelines = {}
        self.lock = threading.Lock()        # one GPU job at a time
        self.voice, self.speed = voice, speed

    def pipeline(self, voice: str):
        from kokoro import KPipeline

        lang = voice[0]
        if lang not in self.pipelines:
            self.pipelines[lang] = KPipeline(lang_code=lang, repo_id=self.REPO, model=self.model)
        return self.pipelines[lang]

    def speak(self, text: str, voice: str | None = None, speed: float | None = None):
        """Audio (float32 numpy) for one passage of text."""
        import numpy as np

        voice, speed = voice or self.voice, speed or self.speed
        with self.lock:
            out = [r.audio.cpu().numpy() if hasattr(r.audio, "cpu") else r.audio
                   for r in self.pipeline(voice)(text, voice=voice, speed=speed, split_pattern=None)
                   if r.audio is not None]
        return np.concatenate(out) if out else np.zeros(0, dtype=np.float32)

    def paragraphs(self, text: str):
        """Yield audio per paragraph, with a short pause after each."""
        import numpy as np

        pause = np.zeros(int(SAMPLE_RATE * 0.35), dtype=np.float32)
        for para in text.split("\n\n"):
            if para.strip():
                yield self.speak(para)
                yield pause

    def chapter_audio(self, ch: Chapter):
        import numpy as np

        title = self.paragraphs(ch.title + ".") if not ch.text.startswith(ch.title) else iter(())
        lead = np.zeros(int(SAMPLE_RATE * 0.5), dtype=np.float32)
        return np.concatenate([lead, *title, *self.paragraphs(ch.text), lead])


# ---------------------------------------------------------------- output

def ffmpeg() -> str:
    import imageio_ffmpeg
    return shutil.which("ffmpeg") or imageio_ffmpeg.get_ffmpeg_exe()


def run(cmd):
    r = subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, text=True, errors="replace")
    if r.returncode:
        raise RuntimeError("ffmpeg failed:\n" + r.stderr[-2000:])


def safe(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip(" .")[:80] or "untitled"


def parse_range(spec: str | None, n: int) -> list[int]:
    if not spec:
        return list(range(n))
    picked = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        lo = int(a) if a else 1
        hi = int(b) if b else (n if _ else lo)
        picked.update(range(lo - 1, min(hi, n)))
    return sorted(i for i in picked if 0 <= i < n)


def book_title(path: Path) -> str:
    if path.suffix.lower() == ".epub":
        try:
            from ebooklib import epub
            meta = epub.read_epub(str(path), options={"ignore_ncx": True}).get_metadata("DC", "title")
            if meta:
                return meta[0][0]
        except Exception:
            pass
    return path.stem


def make_book(path: Path, chapters: list[Chapter], idx: list[int], args, narrator=None, progress=None):
    """Build the audiobook; returns the output path. progress(done, total, title) is called per chapter."""
    import soundfile as sf

    title = book_title(path)
    out_dir = Path(args.out)
    work = out_dir / f".{safe(title)}_work"
    work.mkdir(parents=True, exist_ok=True)
    if narrator is None:
        narrator = Narrator(args.voice, args.speed)
    else:
        narrator.voice, narrator.speed = args.voice, args.speed
    tag = f"{args.voice}_{args.speed}"

    parts = []
    for k, i in enumerate(idx, 1):
        ch = chapters[i]
        part = work / f"{i + 1:03d}_{tag}.m4a"
        if progress:
            progress(k - 1, len(idx), ch.title)
        if not part.exists():               # resumable: finished chapters are kept
            print(f"[{k}/{len(idx)}] {ch.title}  ({len(ch.text):,} chars)")
            audio = narrator.chapter_audio(ch)
            wav = work / "tmp.wav"
            sf.write(wav, audio, SAMPLE_RATE)
            run([ffmpeg(), "-y", "-i", str(wav), "-c:a", "aac", "-b:a", "64k", "-ac", "1", str(part)])
            wav.unlink()
        else:
            print(f"[{k}/{len(idx)}] {ch.title}  (already done)")
        parts.append((ch, part))

    def duration(p: Path) -> float:
        r = subprocess.run([ffmpeg(), "-i", str(p)], stderr=subprocess.PIPE, text=True)
        h, m, s = re.search(r"Duration: (\d+):(\d+):([\d.]+)", r.stderr).groups()
        return int(h) * 3600 + int(m) * 60 + float(s)

    if args.format == "mp3":
        dest = out_dir / safe(title)
        dest.mkdir(parents=True, exist_ok=True)
        for n, (ch, p) in enumerate(parts, 1):
            run([ffmpeg(), "-y", "-i", str(p), "-c:a", "libmp3lame", "-b:a", "96k",
                 "-metadata", f"title={ch.title}", "-metadata", f"album={title}",
                 "-metadata", f"track={n}", str(dest / f"{n:03d} - {safe(ch.title)}.mp3")])
        print(f"\nDone: {dest}")
    else:
        listing = work / "list.txt"
        listing.write_text("".join("file '" + p.resolve().as_posix().replace("'", "'\\''") + "'\n"
                                   for _, p in parts), encoding="utf-8")
        meta = [";FFMETADATA1", f"title={title}", f"album={title}", "genre=Audiobook"]
        t = 0.0
        for ch, p in parts:
            d = duration(p)
            esc = re.sub(r"([=;#\\\n])", r"\\\1", ch.title)
            meta += ["[CHAPTER]", "TIMEBASE=1/1000", f"START={int(t * 1000)}",
                     f"END={int((t + d) * 1000)}", f"title={esc}"]
            t += d
        meta_file = work / "meta.txt"
        meta_file.write_text("\n".join(meta) + "\n", encoding="utf-8")
        dest = out_dir / f"{safe(title)}.m4b"
        run([ffmpeg(), "-y", "-f", "concat", "-safe", "0", "-i", str(listing), "-i", str(meta_file),
             "-map_metadata", "1", "-map_chapters", "1", "-c", "copy", str(dest)])
        print(f"\nDone: {dest}  ({t / 3600:.1f} h)")

    if not args.keep_work:
        shutil.rmtree(work, ignore_errors=True)
    if progress:
        progress(len(idx), len(idx), "Done")
    return dest


def play(chapters: list[Chapter], idx: list[int], args):
    import queue
    import threading
    import sounddevice as sd

    narrator = Narrator(args.voice, args.speed)
    q: queue.Queue = queue.Queue(maxsize=8)

    def produce():
        for i in idx:
            ch = chapters[i]
            q.put(("title", f"Chapter {i + 1}: {ch.title}"))
            if not ch.text.startswith(ch.title):
                for a in narrator.paragraphs(ch.title + "."):
                    q.put(("audio", a))
            for a in narrator.paragraphs(ch.text):
                q.put(("audio", a))
        q.put(("end", None))

    threading.Thread(target=produce, daemon=True).start()
    print("Playing. Ctrl+C to stop.")
    try:
        while True:
            kind, val = q.get()
            if kind == "end":
                break
            if kind == "title":
                print(f"\n>> {val}")
            else:
                sd.play(val, SAMPLE_RATE)
                sd.wait()
    except KeyboardInterrupt:
        sd.stop()
        print("\nStopped.")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Narrate ebooks into audiobooks.")
    ap.add_argument("book", nargs="?", help="EPUB, PDF, DOCX, or TXT file")
    ap.add_argument("--voice", default="af_heart", help="Kokoro voice (see --list-voices)")
    ap.add_argument("--speed", type=float, default=1.0, help="speech speed, e.g. 0.9 or 1.2")
    ap.add_argument("--format", choices=["m4b", "mp3"], default="m4b",
                    help="m4b = one audiobook file with chapters; mp3 = one file per chapter")
    ap.add_argument("--chapters", help="which chapters, e.g. 1-3,7 or 5- (default: all)")
    ap.add_argument("--out", default="output", help="output folder")
    ap.add_argument("--play", action="store_true", help="read aloud live instead of making a file")
    ap.add_argument("--list-chapters", action="store_true")
    ap.add_argument("--list-voices", action="store_true")
    ap.add_argument("--keep-work", action="store_true", help="keep per-chapter intermediate audio")
    args = ap.parse_args()

    if args.list_voices:
        for group, names in VOICES.items():
            print(f"{group:16} {names}")
        return
    if not args.book:
        ap.error("give a book file")
    path = Path(args.book)
    if not path.exists():
        sys.exit(f"Not found: {path}")

    chapters = extract(path)
    if not chapters:
        sys.exit("No readable text found (scanned PDFs need OCR first).")
    idx = parse_range(args.chapters, len(chapters))

    if args.list_chapters:
        for i, ch in enumerate(chapters, 1):
            print(f"{i:4}  {len(ch.text):>9,} chars  {ch.title}")
        total = sum(len(c.text) for c in chapters)
        print(f"\n{len(chapters)} chapters, {total:,} chars (~{total / 900 / 60:.1f} h of audio)")
        return

    if args.play:
        play(chapters, idx, args)
    else:
        make_book(path, chapters, idx, args)


if __name__ == "__main__":
    main()
